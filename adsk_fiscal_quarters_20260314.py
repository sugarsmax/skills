from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

wb = Workbook()
ws = wb.active
ws.title = "Fiscal Quarters"

header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill("solid", start_color="1D3557")
header_align = Alignment(horizontal="center")

data_font = Font(name="Arial")
alt_fill = PatternFill("solid", start_color="EBF4FF")

ws["A1"] = "Quarter"
ws["A1"].font = header_font
ws["A1"].fill = header_fill
ws["A1"].alignment = header_align

row = 2
for fy in range(24, 29):
    for q in range(1, 5):
        ws[f"A{row}"] = f"FY{fy}-Q{q}"
        ws[f"A{row}"].font = data_font
        ws[f"A{row}"].alignment = Alignment(horizontal="center")
        if row % 2 == 0:
            ws[f"A{row}"].fill = alt_fill
        row += 1

ws.column_dimensions["A"].width = 14
ws.auto_filter.ref = ws.dimensions
ws.freeze_panes = "B2"

out = Path.home() / "Downloads" / "adsk_fiscal_quarters_20260314.xlsx"
wb.save(out)
print(f"Saved: {out}")
